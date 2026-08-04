"""Read TXT, PDF and DOCX documents for the local RAG knowledge base."""
from __future__ import annotations

import os
import re
import zipfile


SUPPORTED_EXTS = (".txt", ".pdf", ".docx")


def load_txt(file_path: str) -> str:
    for encoding in ("utf-8", "gbk"):
        try:
            with open(file_path, "r", encoding=encoding) as handle:
                return handle.read()
        except (UnicodeDecodeError, OSError):
            continue
    return ""


def read_pdf_sections(file_path: str) -> list[dict]:
    try:
        from pypdf import PdfReader
    except ImportError:
        from PyPDF2 import PdfReader

    sections = []
    for index, page in enumerate(PdfReader(file_path).pages, start=1):
        text = page.extract_text() or ""
        if text.strip():
            sections.append({"text": text, "page": index})
    return sections


def read_docx(file_path: str) -> str:
    """Read paragraphs, tables and text boxes from a DOCX file."""
    import docx

    document = docx.Document(file_path)
    parts = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    try:
        with zipfile.ZipFile(file_path) as archive:
            xml = archive.read("word/document.xml").decode("utf-8", errors="ignore")
        for block in xml.split("</w:p>"):
            text = "".join(re.findall(r"<w:t[^>]*>([^<]*)</w:t>", block)).strip()
            if text and text not in parts:
                parts.append(text)
    except Exception:
        pass
    return "\n".join(parts)


def read_file_sections(file_path: str) -> list[dict]:
    extension = os.path.splitext(file_path)[1].lower()
    if extension == ".pdf":
        return read_pdf_sections(file_path)
    if extension == ".txt":
        return [{"text": load_txt(file_path), "page": None}]
    if extension == ".docx":
        return [{"text": read_docx(file_path), "page": None}]
    return []


def read_file(file_path: str) -> str:
    return "\n".join(section["text"] for section in read_file_sections(file_path))


def iter_docs(folder: str):
    for root, _dirs, files in os.walk(folder):
        for name in files:
            if name.lower().endswith(SUPPORTED_EXTS):
                yield os.path.join(root, name)
