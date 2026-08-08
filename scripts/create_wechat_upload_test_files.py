from pathlib import Path

from docx import Document
from docx.shared import Pt
from reportlab.lib.pagesizes import A4
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas


OUTPUT_DIR = Path(__file__).resolve().parents[1] / "demo_product_docs" / "wechat_upload_samples"


def create_docx() -> Path:
    path = OUTPUT_DIR / "05_项目立项与采购流程.docx"
    document = Document()
    normal = document.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal.font.size = Pt(10.5)
    document.add_heading("项目立项与采购流程（测试资料）", level=0)
    document.add_paragraph("本文件用于验证 DOCX 上传、解析、检索和引用展示。")
    document.add_heading("一、立项条件", level=1)
    document.add_paragraph("涉及新项目、外包服务或单笔预算超过 5000 元的采购，须先提交项目立项申请，明确目标、预算、负责人和预计完成时间。")
    document.add_heading("二、采购流程", level=1)
    for text in [
        "需求部门提交采购申请和技术规格说明。",
        "采购人员至少收集两家供应商报价，并完成比价记录。",
        "部门负责人和财务分别完成业务与预算审批。",
        "采购合同审批完成后方可下单，付款须附合同和验收记录。",
    ]:
        document.add_paragraph(text, style="List Bullet")
    document.add_heading("三、验收与归档", level=1)
    document.add_paragraph("项目交付后由需求部门填写验收单，采购资料、合同、发票和付款凭证统一归档，保存期限不少于三年。")
    document.save(path)
    return path


def create_pdf() -> Path:
    path = OUTPUT_DIR / "06_差旅与会议管理.pdf"
    font_path = Path(r"C:\Windows\Fonts\msyh.ttc")
    font_name = "Helvetica"
    if font_path.is_file():
        try:
            pdfmetrics.registerFont(TTFont("MicrosoftYaHei", str(font_path), subfontIndex=0))
            font_name = "MicrosoftYaHei"
        except Exception:
            pass
    pdf = canvas.Canvas(str(path), pagesize=A4)
    width, height = A4
    pdf.setFont(font_name, 16)
    pdf.drawString(56, height - 64, "差旅与会议管理（测试资料）")
    pdf.setFont(font_name, 10.5)
    lines = [
        "本文件用于验证 PDF 上传、解析、检索和引用展示。",
        "一、出差申请：员工应至少提前一个工作日提交出差申请，写明地点、事由和预计费用。",
        "二、交通标准：市内优先使用公共交通；跨城出差可选择高铁二等座或普通经济舱。",
        "三、住宿标准：普通员工每晚不超过 350 元，主管每晚不超过 500 元，超标准须提前书面审批。",
        "四、会议费用：会议室、餐饮和物料费用须在会前提交预算，实际支出后五个工作日内报销。",
        "五、资料归档：出差申请、行程凭证、发票和会议签到记录应一并上传到项目资料库。",
    ]
    y = height - 100
    for line in lines:
        pdf.drawString(56, y, line)
        y -= 24
    pdf.save()
    return path


if __name__ == "__main__":
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    print(create_docx())
    print(create_pdf())
