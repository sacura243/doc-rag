from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer


ROOT = Path(__file__).resolve().parents[1] / "demo_product_docs" / "wechat_upload_samples"


def write_text_documents() -> None:
    (ROOT / "01_报销制度.txt").write_text(
        """星河智造（虚拟公司）员工报销制度

一、提交时限
员工应在费用发生后的 3 个工作日内提交报销申请。超过时限需要在申请单中说明原因，并由直属主管确认。

二、必备材料
报销必须提供合规发票或电子发票、报销申请单和费用用途说明。需要审批的费用还要附直属主管的审批记录。差旅费用需要附出差申请和交通、住宿明细。

三、审批与付款
部门负责人在 2 个工作日内完成审批，财务在审批完成后的下一个周五统一付款。单笔超过 2000 元的费用需要财务复核。

四、特殊规则
工作日午餐补贴标准为每人每天 40 元；出差期间住宿标准为普通员工每晚 350 元，主管每晚 500 元。超出标准的部分需要提前获得书面批准。
""",
        encoding="utf-8",
    )
    (ROOT / "02_请假与考勤.txt").write_text(
        """星河智造（虚拟公司）请假与考勤规则

一、请假申请
员工应在休假开始前至少 1 个工作日提交申请。连续请假超过 3 个工作日，需要同时抄送部门负责人和人事专员。

二、病假
病假超过 1 个工作日，需要提交医院或正规互联网医院出具的就诊证明。突发就医无法提前申请时，应在当天上午 10 点前向直属主管报备，并在返岗后补交材料。

三、年假
入职满一年后，每年享有 5 个工作日年假。年假可以按半天申请，未使用年假不能自动折算为工资。

四、考勤
工作日标准出勤时间为 09:00 至 18:00，午休时间为 12:00 至 13:00。迟到超过 30 分钟需要提交说明；一个月内累计迟到 3 次将触发主管沟通。
""",
        encoding="utf-8",
    )


def set_docx_style(style, font_name: str, size: int, color: str, *, bold: bool = False) -> None:
    style.font.name = font_name
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor.from_string(color)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)


def write_docx() -> None:
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

    normal = document.styles["Normal"]
    set_docx_style(normal, "Microsoft YaHei", 11, "263238")
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25
    heading1 = document.styles["Heading 1"]
    set_docx_style(heading1, "Microsoft YaHei", 16, "2E74B5", bold=True)
    heading1.paragraph_format.space_before = Pt(18)
    heading1.paragraph_format.space_after = Pt(10)
    heading2 = document.styles["Heading 2"]
    set_docx_style(heading2, "Microsoft YaHei", 13, "2E74B5", bold=True)
    heading2.paragraph_format.space_before = Pt(14)
    heading2.paragraph_format.space_after = Pt(7)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(4)
    run = title.add_run("星河智造新员工入职指南")
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(22)
    run.font.bold = True
    run.font.color.rgb = RGBColor.from_string("0F766E")
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(16)
    subtitle.add_run("虚拟测试资料 · 适用于知答库上传与引用测试").italic = True

    document.add_heading("入职前三天安排", level=1)
    document.add_heading("第一天：完成基础报到", level=2)
    document.add_paragraph("新员工到达公司后先到人事部报到，领取工牌、电脑和办公用品。人事专员会介绍考勤规则、办公区域和信息安全要求。当天 17:00 前需要完成员工信息确认。")
    document.add_heading("第二天：完成系统与安全培训", level=2)
    document.add_paragraph("第二天上午参加企业微信、代码仓库和内部知识库培训；下午参加信息安全培训，完成密码设置和设备使用确认。安全培训通过后才能申请访问项目资料。")
    document.add_heading("第三天：进入团队工作", level=2)
    document.add_paragraph("第三天由直属导师带领熟悉团队流程，阅读岗位相关制度，并完成一次小型工作任务。导师需要在当天 18:00 前提交入职跟进记录。")
    document.add_heading("入职联系人", level=1)
    document.add_paragraph("人事专员：林晓雨；直属导师：周明远；IT 支持邮箱：it-support@example.test。以上联系人均为虚拟测试信息，不对应真实个人或公司。")

    document.save(ROOT / "03_新员工入职指南.docx")


def write_pdf() -> None:
    font_path = Path("C:/Windows/Fonts/simhei.ttf")
    pdfmetrics.registerFont(TTFont("SimHei", str(font_path)))
    output = ROOT / "04_IT设备使用规范.pdf"
    document = SimpleDocTemplate(
        str(output), pagesize=A4, leftMargin=22 * mm, rightMargin=22 * mm,
        topMargin=20 * mm, bottomMargin=20 * mm,
    )
    styles = getSampleStyleSheet()
    title = ParagraphStyle("TestTitle", parent=styles["Title"], fontName="SimHei", fontSize=20, leading=27, alignment=TA_CENTER, textColor="#0F766E", spaceAfter=16)
    heading = ParagraphStyle("TestHeading", parent=styles["Heading2"], fontName="SimHei", fontSize=13, leading=19, textColor="#2E74B5", spaceBefore=10, spaceAfter=6)
    body = ParagraphStyle("TestBody", parent=styles["BodyText"], fontName="SimHei", fontSize=10.5, leading=17, textColor="#263238", spaceAfter=7)
    story = [Paragraph("星河智造 IT 设备使用规范", title), Paragraph("虚拟测试资料 · 用于测试 PDF 上传、解析和引用来源", body)]
    sections = [
        ("一、设备领用", "员工领取电脑、显示器或移动设备时，需要核对设备编号并在资产登记表上签字。设备仅限本人工作使用，不得转借给未授权人员。"),
        ("二、密码要求", "公司系统密码长度不得少于 12 位，必须同时包含大小写字母、数字和特殊字符。密码每 90 天更换一次，不能使用姓名、生日或连续数字。"),
        ("三、远程办公", "远程办公必须通过公司 VPN 访问内部系统。公共场所使用电脑时应开启屏幕锁定，离开座位超过 5 分钟必须锁屏。"),
        ("四、故障报修", "设备出现故障时，先记录设备编号、故障现象和发生时间，再发送至 it-support@example.test。IT 支持工作日 2 小时内响应，紧急故障优先处理。"),
    ]
    for heading_text, paragraph in sections:
        story.extend([Paragraph(heading_text, heading), Paragraph(paragraph, body)])
    document.build(story)


if __name__ == "__main__":
    ROOT.mkdir(parents=True, exist_ok=True)
    write_text_documents()
    write_docx()
    write_pdf()
    print(ROOT)
